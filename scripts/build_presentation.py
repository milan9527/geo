#!/usr/bin/env python3
"""Create the English Aperture deck from real screenshots and generated diagrams."""
import json
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs/presentation"
SHOTS = ROOT / "docs/screenshots"
DIAGRAMS = ROOT / "docs/architecture"
W, H = 13.333333, 7.5
NAVY, TEAL, INK, MUTED, BG = "122330", "0E9384", "152A36", "56697A", "F4F7F9"
MINT, WHITE, LINE = "67DCCC", "FFFFFF", "D9E2E8"
DATE = "25 September 2026"
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
prs.core_properties.title = "Aperture Intelligence — English Technical Overview"
prs.core_properties.subject = "Evidence-based research, bilingual publication and agent access on AWS"
prs.core_properties.author = "Aperture Intelligence"
prs.core_properties.keywords = "AWS, AgentCore, bilingual publishing, editorial review, x402"
prs.core_properties.comments = "Real English production screenshots. Dated verification; testnet payments."
manifest = json.loads((SHOTS / "manifest.json").read_text())
shots = {item["file"]: item for item in manifest["screenshots"]}
slide_manifest = []


def color(value):
    return RGBColor.from_string(value)


def box(slide, x, y, w, h, fill=WHITE, line=None, radius=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                  Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    if line:
        shape.line.color.rgb = color(line)
        shape.line.width = Pt(.6)
    else:
        shape.line.fill.background()
    shape._element.spPr.append(OxmlElement("a:effectLst"))
    if radius:
        shape.adjustments[0] = .06
    return shape


def text(slide, x, y, w, h, value, size=20, fill=INK, bold=False, font="Arial"):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.TOP
    for i, line in enumerate(value.split("\n")):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = line
        p.font.name, p.font.size, p.font.bold = font, Pt(size), bold
        p.font.color.rgb = color(fill)
        p.space_after = Pt(8 if size >= 18 else 3)
        p.line_spacing = 1.12
    return shape


def new(title, section, dark=False, notes=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color(NAVY if dark else BG)
    slide_manifest.append({"slide": len(prs.slides), "title": title, "screenshots": []})
    text(slide, .5, .3, 11.9, .3, section.upper(), 11, MINT if dark else TEAL, True)
    text(slide, .5, .83, 12.3, .8, title, 32, WHITE if dark else INK, True)
    footer(slide, dark)
    slide.notes_slide.notes_text_frame.text = notes
    return slide


def footer(slide, dark=False):
    text(slide, .5, 7.12, 10, .2, "APERTURE INTELLIGENCE  /  ENGLISH EDITION  /  25 SEP 2026",
         9, "9CADB8" if dark else MUTED)
    text(slide, 12.1, 7.08, .7, .27, f"{len(prs.slides):02d}", 11, MINT if dark else TEAL, True)


def image(slide, path, x, y, w, h, border=True):
    iw, ih = Image.open(path).size
    factor = min(w / iw, h / ih)
    dw, dh = iw * factor, ih * factor
    px, py = x + (w - dw) / 2, y + (h - dh) / 2
    if border:
        box(slide, px - .025, py - .025, dw + .05, dh + .05, WHITE, LINE)
    pic = slide.shapes.add_picture(str(path), Inches(px), Inches(py), width=Inches(dw), height=Inches(dh))
    pic.name = path.name
    if path.name in shots:
        item = shots[path.name]
        slide_manifest[-1]["screenshots"].append(path.name)
        current = slide.notes_slide.notes_text_frame.text
        slide.notes_slide.notes_text_frame.text = (
            current + f"\nScreenshot: {item['caption']}\nSource: {item['url']}\n"
            f"Captured: {item['capturedAt']}\nSHA-256: {item['sha256']}\n"
            f"Region: {item['region']}. Actual English interface, no substituted text or synthetic data."
        )
    return pic


def callout(slide, x, y, w, number, heading, body, dark=False):
    fill = "1B3443" if dark else WHITE
    box(slide, x, y, w, 3.9, fill, None if dark else LINE, radius=True)
    text(slide, x + .25, y + .25, w - .5, .7, number, 35, MINT if dark else TEAL, True)
    text(slide, x + .25, y + 1.1, w - .5, .9, heading, 23, WHITE if dark else INK, True)
    text(slide, x + .25, y + 2.03, w - .5, 1.55, body, 19, "CDD9E0" if dark else MUTED)


def diagram(slug, title, note):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color(BG)
    slide_manifest.append({"slide": len(prs.slides), "title": title, "screenshots": []})
    slide.notes_slide.notes_text_frame.text = note + f"\nEditable source: docs/architecture/aperture-aws.drawio, page {title}."
    image(slide, DIAGRAMS / f"{slug}.png", .08, .02, W - .16, 7.15, border=False)
    footer(slide)
    return slide


def main():
    slide = new("Aperture Intelligence", "AWS reference application", dark=True,
                notes="Technical review and project demonstration. Snapshot date: 2026-09-25. "
                      "The application connects research production to stable bilingual delivery, search and agent access.")
    text(slide, .52, 2.0, 4.45, 1.55, "From public evidence\nto reviewed research.", 31, WHITE, True)
    text(slide, .52, 4.12, 4.35, 1.45,
         "English-first publishing.\nPermanent public URLs.\nProgrammatic agent access.", 21, "CCDCE4")
    text(slide, .52, 6.38, 5, .35, "Technical overview  •  " + DATE, 13, MINT)
    image(slide, SHOTS / "public-home-en.png", 5.3, 1.87, 7.48, 4.68)

    slide = new("One system, three useful outcomes", "Purpose",
                notes="Purpose statements describe the implemented workflow. Citation, indexing and revenue are outcomes to measure, "
                      "not guarantees. The x402 deployment is a Base Sepolia testnet demonstration.")
    callout(slide, .5, 2.03, 3.9, "01", "Research with evidence",
            "Collect public material, save source evidence, and review generated conclusions.")
    callout(slide, 4.72, 2.03, 3.9, "02", "Publish for discovery",
            "Serve readable English and Chinese editions at stable, crawlable URLs.")
    callout(slide, 8.94, 2.03, 3.9, "03", "Serve people and agents",
            "Offer public HTML, open JSON and an x402 payment-protocol demonstration.")
    text(slide, .55, 6.36, 12.2, .4, "Built for engineers exploring the full research-to-publication lifecycle.", 20, MUTED)

    slide = new("A browsable research publication", "Public experience",
                notes="Real production category page. Full category lists come from SSR and are not capped at three cards or the home API default.")
    image(slide, SHOTS / "public-category-en.png", .5, 1.92, 8.48, 4.92)
    text(slide, 9.4, 2.15, 3.3, .65, "Five research areas", 24, INK, True)
    text(slide, 9.4, 3.08, 3.3, 2.35,
         "AI industry\nAgent technology\nCloud computing\nCommerce and media\nFinancial markets", 20, MUTED)
    text(slide, 9.4, 5.78, 3.25, .8, "Complete categories,\nserver-rendered links.", 18, TEAL, True)

    slide = new("English by default, Chinese alongside", "Bilingual delivery",
                notes="41 approved English editions existed in the September 25 audit. Chinese drafts and original evidence are preserved. "
                      "English and Chinese refer to the same article ID. See docs/bilingual-publication.md.")
    image(slide, SHOTS / "public-article-en.png", 4.65, 1.94, 8.18, 4.87)
    text(slide, .55, 2.12, 3.65, .7, "One article, two editions", 24, INK, True)
    text(slide, .55, 3.15, 3.85, 1.4, "EN  /article/{slug}\nZH  /zh/article/{slug}", 18, TEAL, True)
    text(slide, .55, 4.54, 3.75, 1.87,
         "Full reviewed translations.\nPage-preserving switching.\nCanonical + hreflang.\nBilingual sitemap and RSS.", 20, MUTED)

    slide = new("Discovery continues inside the site", "Search and navigation",
                notes="Actual English search overlay opened through normal UI controls. The query was Agent. "
                      "Navigation preserves metadata, language and browser history; temporary failures are distinguished from real 404s.")
    image(slide, SHOTS / "public-search-en.png", .55, 1.93, 4.57, 4.93)
    text(slide, 5.82, 2.25, 6.65, .7, "Search the reviewed catalogue", 27, INK, True)
    text(slide, 5.82, 3.35, 6.3, 2.2,
         "Search results use the selected language.\nCategories expose every published article.\nRelated reading and RSS offer a next step.\nDirect visits and client navigation share SSR.", 21, MUTED)
    box(slide, 5.8, 5.96, 6.47, .69, "E3F3EF", radius=True)
    text(slide, 6.02, 6.14, 6.08, .34, "Verified discovery paths. No ranking guarantee.", 16, TEAL, True)

    diagram("aws-overview", "Aperture on AWS",
            "Logical component/data flow, not a VPC/subnet topology. The indexing worker is invoked by the publisher/API after a successful commit; "
            "Aurora does not directly invoke Lambda. Read docs/architecture.md for security and persistence boundaries.")

    slide = new("Research runs asynchronously", "Implementation",
                notes="Sources: aws_scheduler/lambda_function.py, aws_runtime/app.py, aws_runtime/crawler_tools.py, agent_runtime/codex_crawler.ts. "
                      "The runtime uses Bedrock. Scheduler delivery retries and Lambda async retries are separate policies.")
    callout(slide, .5, 2.02, 3.9, "1", "Schedule and dispatch",
            "EventBridge schedules a Lambda bridge. The bridge waits for runtime acceptance.")
    callout(slide, 4.72, 2.02, 3.9, "2", "Collect and analyze",
            "The runtime selects sources, generates crawler code and retrieves evidence with managed tools.")
    callout(slide, 8.94, 2.02, 3.9, "3", "Record and review",
            "Aurora stores job state, evidence, research output and publication decisions.")
    text(slide, .55, 6.38, 12.1, .4, "Codex SDK generates crawler code; Browser and Code Interpreter perform retrieval and execution.", 18, MUTED)

    diagram("publication-pipeline", "Review before every new publication",
            "Quality and full published-catalogue semantic comparison precede translation. Independent English review requires score >=90 and no issues. "
            "A final transaction rechecks source/article hashes and the published catalogue. Semantic review is not a mathematical uniqueness guarantee. "
            "Source: aws_runtime/publication.py and docs/bilingual-publication.md.")

    slide = new("New content must not erase old URLs", "Publication and search", dark=True,
                notes="September 25 verification snapshot: 41 articles, 104 bilingual public URLs, 19 legacy redirect mappings. "
                      "312 crawler page checks and 26 site checks passed. Bing accepted 104 URLs at 15:29:24 UTC with HTTP 200. "
                      "The Bing Webmaster indexed-page report was not accessed; Google does not use IndexNow.")
    for x, count, label in [(0.55, "41", "published articles"), (4.78, "104", "bilingual public URLs"),
                            (9.0, "19", "legacy redirect mappings")]:
        text(slide, x, 2.03, 3.7, 1.0, count, 54, MINT, True)
        text(slide, x, 3.12, 3.7, .55, label, 20, WHITE)
    text(slide, .55, 4.27, 12, 1.9,
         "Published status and slugs are protected by the database.\nReviewed repairs use a specific 301 or restore the original URL.\nNew duplicates remain under review; existing pages stay available.", 24, WHITE)
    text(slide, .55, 6.54, 12.05, .3, "Snapshot: 25 Sep 2026  •  Bing accepted the update; acceptance does not confirm indexing.", 14, MINT)

    slide = new("An English console for daily operations", "Administration",
                notes="Actual production overview, captured in an 1600x850 viewport. Historical original-language logs below the viewport were not edited. "
                      "Displayed USDC values are testnet activity and must not be described as production revenue.")
    image(slide, SHOTS / "admin-dashboard-en.png", .5, 1.84, 12.33, 4.96)
    text(slide, .65, 6.84, 11.95, .24, "Traffic estimates, reader attribution, research activity and x402 testnet events in one view.", 13, MUTED)

    slide = new("Published content uses reviewed English editions", "Content operations",
                notes="Actual content view with the normal Published filter; original unpublished drafts are not translated through DOM substitutions. "
                      "The second screenshot is the actual reviewed English article detail modal. Content mutations were not performed.")
    image(slide, SHOTS / "admin-content-en.png", .52, 2.02, 7.36, 4.6)
    image(slide, SHOTS / "admin-article-en.png", 8.25, 2.02, 4.56, 4.6)
    text(slide, .55, 1.61, 7.2, .28, "PUBLISHED CATALOGUE", 11, TEAL, True)
    text(slide, 8.28, 1.61, 4.4, .28, "REVIEWED ARTICLE DETAIL", 11, TEAL, True)
    text(slide, .55, 6.76, 12.1, .29, "Original drafts and source evidence retain their original language and audit history.", 14, MUTED)

    slide = new("Sources and schedules remain inspectable", "Research operations",
                notes="Actual English crawler and source registry screens. Source counts are operational data and can change. "
                      "Settings and source edits were not saved during capture. Source credentials are never embedded in the presentation.")
    image(slide, SHOTS / "admin-crawlers-en.png", .5, 2.05, 6.03, 3.77)
    image(slide, SHOTS / "admin-sources-en.png", 6.8, 2.05, 6.03, 3.77)
    text(slide, .58, 6.12, 5.83, .65, "Crawler state, UTC schedules\nand recorded job outcomes.", 19, MUTED)
    text(slide, 6.88, 6.12, 5.83, .65, "Registered sources, assignments\nand connection-test workflows.", 19, MUTED)

    diagram("x402-payment-flow", "Agent access with x402",
            "Seller implementation: backend/x402_payment.py. Payments are verified and settled by an external facilitator. "
            "The optional research buyer uses AgentCore Payments separately. Base Sepolia testnet USDC only; default 0.002. "
            "No payment was executed while capturing or producing this presentation. Successful response includes content and PAYMENT-RESPONSE.")

    slide = new("Measure observed behavior and expose limits", "Analytics and controls",
                notes="CloudFront -> S3 -> SQS -> Lambda -> Aurora aggregates. HLL and browser events are approximate. "
                      "S3 raw logs have a lifecycle policy; automatic Aurora analytics cleanup is not configured. "
                      "Settings switches persist preferences, not runtime controls. See docs/functional-verification.md.")
    image(slide, SHOTS / "admin-settings-en.png", 6.0, 2.03, 6.82, 4.27)
    text(slide, .55, 2.03, 5.05, .7, "Two measurement sources", 25, INK, True)
    text(slide, .55, 3.03, 5.05, 1.9,
         "Edge logs: requests and visitor estimates.\nBrowser events: sessions, reading,\nreturn visits and referral channels.", 19, MUTED)
    text(slide, .55, 5.24, 5.1, 1.35,
         "Exclude diagnostics from growth.\nRSS clicks do not prove subscription.\nTestnet payments are not revenue.", 19, TEAL, True)
    text(slide, 6.1, 6.45, 6.6, .42, "Preference-only controls are labeled in the actual UI.", 15, MUTED)

    slide = new("A small application with explicit boundaries", "Code organization",
                notes="Python HTTP/SSR backend with vanilla frontend assets; managed AWS research tools and a TypeScript Codex worker. "
                      "Aurora Data API reads are bounded with pagination/chunks and consistent snapshots for large admin responses. "
                      "No separate graph/vector database is part of the implemented architecture.")
    rows = [
        ("frontend/public + frontend/admin", "Reader experience and operations console"),
        ("backend", "SSR, APIs, auth, locales, persistence and x402 seller"),
        ("aws_runtime + agent_runtime", "Research orchestration, tools and publication gates"),
        ("aws_scheduler + aws_indexing_notifier", "Asynchronous invocation and discovery notifications"),
        ("aws_traffic_aggregator", "Edge-log classification and hourly aggregates"),
        ("scripts + infrastructure/aws", "Release workflows, tests, maintenance and IAM examples"),
    ]
    for i, (path, purpose) in enumerate(rows):
        y = 1.98 + i * .72
        box(slide, .52, y, 12.28, .62, WHITE if i % 2 == 0 else "EAF0F3")
        text(slide, .73, y + .13, 5.3, .38, path, 17, TEAL, True)
        text(slide, 6.03, y + .13, 6.53, .38, purpose, 17, INK)
    text(slide, .57, 6.57, 12.0, .4, "Large Data API responses use bounded reads; review history and evidence remain intact.", 19, MUTED)

    slide = new("Run the application locally", "Install",
                notes="Source: scripts/start.sh, scripts/create_admin_user.py and docs/deployment.md. "
                      "Startup creates a venv, installs dependencies and starts local PostgreSQL. Do not load production credentials for local tests. "
                      "The admin image is the actual production English login screen, included as UI illustration, not a claim of a local capture.")
    text(slide, .55, 1.96, 7.25, .65, "Python 3.10+  •  Docker Compose", 23, INK, True)
    box(slide, .52, 2.79, 7.37, 1.38, NAVY, radius=True)
    text(slide, .76, 3.01, 6.9, 1.08,
         "git clone https://github.com/milan9527/geo.git\ncd geo\n./scripts/start.sh", 16, "E1F7F1", font="Liberation Mono")
    text(slide, .55, 4.49, 7.13, 1.16,
         "Public :4173  /  Admin :4174  /  API :8000\nCreate an admin with scripts/create_admin_user.py.\nThe command prompts for a password.", 19, MUTED)
    text(slide, .55, 6.11, 7.0, .64, "Node.js 24+ is needed for the crawler worker build.\nSee docs/deployment.md for configuration.", 16, TEAL, True)
    image(slide, SHOTS / "admin-login-en.png", 8.26, 2.28, 4.55, 3.9)
    text(slide, 8.33, 6.16, 4.36, .35, "Actual English production login screen", 12, MUTED)

    slide = new("Release into an existing AWS environment", "Deploy", dark=True,
                notes="Source: docs/deployment.md, infrastructure/aws/README.md and scripts/deploy_*.sh. "
                      "These scripts are not complete new-account infrastructure bootstraps. Account-specific policy/scheduler files must be adapted. "
                      "Do not delete/update the legacy CloudFormation stack that still owns live resources. Never run mutating smoke tests against production.")
    text(slide, .55, 1.94, 11.95, 1.02,
         "Export .env.aws and .env.deploy.aws with your resource settings.\nProvision the required AWS services before running release scripts.", 22, "D5E3E9")
    box(slide, .54, 3.23, 12.23, 1.85, "1B3443", radius=True)
    text(slide, .83, 3.48, 11.55, 1.42,
         "./scripts/deploy_web_ecs.sh\n./scripts/deploy_agent_runtime.sh --auto-publish\n./scripts/deploy_scheduler_bridge.sh", 22, MINT, font="Liberation Mono")
    text(slide, .55, 5.47, 12.12, 1.35,
         "Build and scan images → update services → upload assets → invalidate caches.\nWait for deployment and cache completion, then verify live bilingual routes.\nKeep isolated tests separate from the production database.", 21, WHITE)

    slide = new("Verified behavior, with evidence you can inspect", "Verification",
                notes="Sources: reports/functional-coverage-2026-09-25.md and JSON; docs/screenshots/manifest.json. "
                      "106 tests and 312 crawler checks were completed before this documentation task. Screenshot captures are additional read-only evidence. "
                      "Known uncached category latency and preference-only controls are documented. Test results do not establish search ranking or adoption.")
    for x, number, label in [(.55, "106", "automated tests passed"),
                              (4.75, "312", "crawler page checks passed"),
                              (8.95, "11", "real English UI captures")]:
        text(slide, x, 2.0, 3.75, .9, number, 49, TEAL, True)
        text(slide, x, 3.05, 3.75, .55, label, 18, MUTED)
    text(slide, .55, 4.03, 11.98, 1.25,
         "Both languages. Publication gates. Navigation and admin workflows.\nLegacy redirects. Payment responses. Traffic attribution.", 23, INK, True)
    text(slide, .55, 5.66, 11.97, 1.07,
         "Explore the site     aperture.zhangwangshu.com\nRead the code        github.com/milan9527/geo\nReproduce checks     docs/functional-verification.md", 19, MUTED)
    for slide in prs.slides:
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= prs.slide_width + 100
            assert shape.top + shape.height <= prs.slide_height + 100
    OUT.mkdir(parents=True, exist_ok=True)
    destination = OUT / "aperture-intelligence-en.pptx"
    prs.save(destination)
    (OUT / "manifest.json").write_text(json.dumps({
        "title": prs.core_properties.title, "language": "en", "snapshotDate": "2026-09-25",
        "slides": slide_manifest, "sourceScreenshotManifest": "../screenshots/manifest.json",
        "notes": "Native editable slide text; diagrams also supplied as editable draw.io. Screenshots are actual UI captures.",
    }, indent=2) + "\n")
    print(f"{destination}: {len(prs.slides)} slides")


if __name__ == "__main__":
    main()
