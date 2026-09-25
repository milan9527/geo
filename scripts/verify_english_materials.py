#!/usr/bin/env python3
"""Verify documentation artifacts without contacting production services."""
import hashlib
import json
from pathlib import Path
import re
import subprocess
from urllib.parse import unquote
import xml.etree.ElementTree as ET
from zipfile import ZipFile

import pymupdf
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
CJK = re.compile(r"[\u3400-\u9fff]")


def main():
    deck = ROOT / "docs/presentation/aperture-intelligence-en.pptx"
    pdf_path = deck.with_suffix(".pdf")
    prs = Presentation(deck)
    pdf = pymupdf.open(pdf_path)
    assert len(prs.slides) == len(pdf) == 18
    pdf[0].get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5)).save(str(deck.parent / "preview.png"))
    paths = subprocess.check_output(
        ["git", "ls-files", "--cached", "--others", "--exclude-standard", "-z"],
        cwd=ROOT,
    ).decode().split("\0")
    markdown = sorted({ROOT / p for p in paths if p.endswith(".md")})
    reference_drafts = {
        "juejin-data_api_1mb.md", "juejin-permanent_urls.md",
        "zhihu-data-api.md", "zhihu-permanent-urls.md",
    }
    checked_links = 0
    for path in markdown:
        content = path.read_text()
        is_source = (path.parent == ROOT / "content/distribution/2026-09-20"
                     and path.name in reference_drafts)
        if not is_source:
            assert not CJK.search(content), f"Non-English project prose: {path}"
        for target in re.findall(r"\]\(([^)\n]+)\)", content):
            target = target.split('"')[0].strip().strip("<>")
            if re.match(r"^[a-z]+:", target) or target.startswith("#"):
                continue
            local = unquote(target.split("#")[0])
            assert (path.parent / local).exists(), f"Missing link: {path} -> {target}"
            checked_links += 1

    drawio = ET.parse(ROOT / "docs/architecture/aperture-aws.drawio").getroot()
    architecture = ROOT / "docs/architecture"
    exports = json.loads((architecture / "manifest.json").read_text())
    assert exports["sourceSha256"] == hashlib.sha256(
        (architecture / "aperture-aws.drawio").read_bytes()).hexdigest()
    assert len(drawio) == 3
    diagram_labels = 0
    native_aws_icons = 0
    diagram_hashes = set()
    for diagram in drawio:
        assert diagram.find("mxGraphModel/root") is not None
        icons = 0
        for cell in diagram.findall(".//mxCell"):
            value = cell.get("value", "")
            assert not CJK.search(value), value
            diagram_labels += bool(value)
            style = cell.get("style", "")
            if "shape=mxgraph.aws4.resourceIcon;" in style:
                assert "resIcon=mxgraph.aws4." in style
                assert "strokeColor=#FFFFFF;" in style
                icons += 1
        export = next(item for item in exports["pages"] if item["page"] == diagram.get("id"))
        assert icons == export["nativeAwsIcons"] and icons > 0
        assert export["textOverflows"] == 0
        native_aws_icons += icons
        for suffix in (".png", ".svg"):
            artifact = architecture / (diagram.get("id") + suffix)
            digest = hashlib.sha256(artifact.read_bytes()).hexdigest()
            assert digest == export[suffix[1:] + "Sha256"], artifact
            if suffix == ".png":
                diagram_hashes.add(digest)

    capture = json.loads((ROOT / "docs/screenshots/manifest.json").read_text())
    assert capture["temporaryAdminRemoved"] and not capture["browserErrors"]
    assert not capture["contentMutations"] and not capture["paymentsExecuted"]
    screenshot_hashes = set()
    for shot in capture["screenshots"]:
        value = hashlib.sha256((ROOT / "docs/screenshots" / shot["file"]).read_bytes()).hexdigest()
        assert value == shot["sha256"], shot["file"]
        screenshot_hashes.add(value)
    assert len(screenshot_hashes) == 11

    with ZipFile(deck) as archive:
        assert archive.testzip() is None
        embedded = {
            hashlib.sha256(archive.read(name)).hexdigest()
            for name in archive.namelist() if name.startswith("ppt/media/")
        }
        assert screenshot_hashes <= embedded, "Missing or modified embedded screenshot"
        assert diagram_hashes <= embedded, "PPTX diagrams do not match the current draw.io exports"
    native_text_boxes = 0
    for index, slide in enumerate(prs.slides):
        rendered = re.sub(r"\s+", "", pdf[index].get_text())
        assert slide.notes_slide.notes_text_frame.text.strip(), f"Missing notes: {index + 1}"
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= prs.slide_width + 100
            assert shape.top + shape.height <= prs.slide_height + 100
            if shape.has_text_frame and shape.text:
                assert not CJK.search(shape.text), shape.text
                expected = re.sub(r"\s+", "", shape.text)
                assert expected in rendered, f"Missing rendered slide text {index + 1}: {shape.text}"
                native_text_boxes += 1
    result = {
        "englishProjectMarkdownFiles": len(markdown) - len(reference_drafts),
        "preservedChineseEditorialDrafts": len(reference_drafts),
        "localDocumentationLinksChecked": checked_links,
        "editableDiagramPages": len(drawio),
        "englishDiagramLabels": diagram_labels,
        "nativeAwsServiceIcons": native_aws_icons,
        "diagramExportsMatchDrawioSource": True,
        "pptxDiagramsMatchCurrentExports": True,
        "realEnglishScreenshots": len(screenshot_hashes),
        "screenshotHashesMatch": True,
        "allScreenshotsEmbeddedUnmodified": True,
        "powerPointSlides": len(prs.slides),
        "renderedPdfPages": len(pdf),
        "nativeTextBoxesFoundInRenderedPdf": native_text_boxes,
        "slideShapesWithinCanvas": True,
        "temporaryCaptureAdminRemoved": True,
        "productionContentMutations": False,
        "paymentsExecuted": False,
        "validationScope": "Documentation artifacts only; earlier application results remain dated in their reports.",
    }
    (deck.parent / "verification.json").write_text(json.dumps(result, indent=2) + "\n")
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
